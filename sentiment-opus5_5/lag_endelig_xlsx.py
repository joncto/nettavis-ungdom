#!/usr/bin/env python3
"""
Lag den endelige leveransefila for kollegaer fra resultatet av trinn 6.

Bruk: uv run lag_endelig_xlsx.py

Leser (kun lesing, endrer ingenting):
    sentiment-ungdom-trinn6-validering.jsonl   (conc_id, dhlabid, concordance, sentiment, label, begrunnelse)
Skriver:
    nettavis-ungdom-sentiment-opus5_5.xlsx

Kolonner: conc_id, dhlabid, konkordanse, nøkkelord, sentiment, label, begrunnelse
- «konkordanse» er teksten uten <b>-tagger, der nøkkelordet vises i fet skrift.
- «nøkkelord» er de uthevede ordene, unike og i rekkefølge, skilt med komma.
Formatet følger eksempel-trinn4-resultater-sammenlimt.xlsx.
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font

ROT = Path(__file__).resolve().parent
KILDE = ROT / "sentiment-ungdom-trinn6-validering.jsonl"
UT = ROT / "nettavis-ungdom-sentiment-opus5_5.xlsx"

KOLONNER = ["conc_id", "dhlabid", "konkordanse", "nøkkelord", "sentiment", "label", "begrunnelse"]
BREDDER = {"conc_id": 10, "dhlabid": 12, "konkordanse": 95, "nøkkelord": 14,
           "sentiment": 14, "label": 14, "begrunnelse": 55}
SENTIMENT = ["positiv", "negativ", "nøytral", "uklart", "ikke relevant"]

FONT = Font(name="Arial", size=10)
FONT_FET = Font(name="Arial", size=10, bold=True)
VANLIG = InlineFont(rFont="Arial", sz=10, b=False)
FET = InlineFont(rFont="Arial", sz=10, b=True)
TOPP = Alignment(vertical="top")
TOPP_BRYT = Alignment(vertical="top", wrap_text=True)
DEL = re.compile(r"(<b>.*?</b>)", re.S)


def konkordanse_og_nokkelord(tekst):
    """Del opp på <b>…</b>; returner (rik tekst, nøkkelord-streng). Andre tegn beholdes uendret."""
    blokker, nokkel = [], []
    for bit in DEL.split(tekst):
        if not bit:
            continue
        if bit.startswith("<b>") and bit.endswith("</b>"):
            ord_ = bit[3:-4]
            blokker.append(TextBlock(FET, ord_))
            if ord_ not in nokkel:
                nokkel.append(ord_)
        else:
            blokker.append(TextBlock(VANLIG, bit))
    return CellRichText(blokker), ", ".join(nokkel)


def main():
    rader = [json.loads(l) for l in open(KILDE, encoding="utf-8") if l.strip()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Konkordanser"
    ws.append(KOLONNER)
    for c in ws[1]:
        c.font = FONT_FET
        c.alignment = TOPP_BRYT

    for r in rader:
        rik, nokkel = konkordanse_og_nokkelord(r["concordance"])
        ws.append([r["conc_id"], r["dhlabid"], rik, nokkel, r["sentiment"], r["label"], r["begrunnelse"]])

    for i, k in enumerate(KOLONNER, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = BREDDER[k]
    for rad in ws.iter_rows(min_row=2):
        for c in rad:
            c.font = FONT
            c.alignment = TOPP_BRYT if c.column in (3, 7) else TOPP
        rad[0].number_format = "@"                       # conc_id som tekst (ledende nuller)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"

    # Sammendrag med formler, som i eksempelfila
    s = wb.create_sheet("Sammendrag")
    n = ws.max_row
    s["A1"] = "Fordeling – Opus 5.5 (claude-opus-5-5), alle konkordanser"
    s["A1"].font = Font(name="Arial", size=11, bold=True)
    s.append(["sentiment", "antall", "andel"])
    for c in s[2]:
        c.font = FONT_FET
    for i, snt in enumerate(SENTIMENT, 3):
        s[f"A{i}"] = snt
        s[f"B{i}"] = f"=COUNTIF(Konkordanser!$E$2:$E${n},$A{i})"
        s[f"C{i}"] = f"=IFERROR($B{i}/$B$8,0)"
    s["A8"], s["B8"], s["C8"] = "Sum", "=SUM($B$3:$B$7)", "=SUM(C3:C7)"
    s["A10"] = "label=ungdomsorg"
    s["B10"] = f'=COUNTIF(Konkordanser!$F$2:$F${n},"ungdomsorg")'
    s["C10"] = "=IFERROR($B10/$B$8,0)"
    s["A12"] = ("Kilde: sentiment-ungdom-trinn6-validering.jsonl (trinn 6). Identiske konkordansetekster er "
                "vurdert én gang og vurderingen kopiert til alle rader; se duplikater-trinn6.xlsx.")
    for rad in s.iter_rows(min_row=3):
        for c in rad:
            if c.value is not None:
                c.font = FONT
    for i in (3, 4, 5, 6, 7, 8, 10):
        s[f"C{i}"].number_format = "0.0%"
    s.column_dimensions["A"].width = 25
    s.column_dimensions["B"].width = 12
    s.column_dimensions["C"].width = 10

    wb.save(UT)
    print(f"Skrev {len(rader)} rader til {UT.name}")


if __name__ == "__main__":
    main()
